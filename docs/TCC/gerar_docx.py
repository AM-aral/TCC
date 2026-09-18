# -*- coding: utf-8 -*-
"""
Gera o documento DOCX do TCC a partir do arquivo TCC_LFG.md.

Uso:
    python gerar_docx.py

Saída:
    TCC_LFG.docx
"""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MD_PATH = "TCC_LFG.md"
OUT_PATH = "TCC_LFG.docx"
FONTE = "Times New Roman"

PRETO = RGBColor(0, 0, 0)


def estilo_fonte(run, tamanho=12, negrito=False, italico=False, fonte=FONTE):
    run.font.name = fonte
    run.font.size = Pt(tamanho)
    run.bold = negrito
    run.italic = italico
    run.font.color.rgb = PRETO
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), fonte)
    rfonts.set(qn("w:hAnsi"), fonte)
    rfonts.set(qn("w:eastAsia"), fonte)


def adicionar_runs_estilizados(paragrafo, texto, tamanho=12, negrito=False, italico=False):
    """Renderiza **negrito**, *italico* e `codigo` inline."""
    tokens = re.split(r"(\*\*.*?\*\*|\*.*?\*|`[^`]*`)", texto)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragrafo.add_run(token[2:-2])
            estilo_fonte(run, tamanho=tamanho, negrito=True, italico=italico)
        elif token.startswith("`") and token.endswith("`"):
            run = paragrafo.add_run(token[1:-1])
            estilo_fonte(run, tamanho=tamanho, fonte="Consolas", italico=italico)
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            run = paragrafo.add_run(token[1:-1])
            estilo_fonte(run, tamanho=tamanho, negrito=negrito, italico=True)
        else:
            run = paragrafo.add_run(token)
            estilo_fonte(run, tamanho=tamanho, negrito=negrito, italico=italico)


def novo_paragrafo(doc, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = alinhamento
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    return p


def paragrafo_corpo(doc, texto):
    p = novo_paragrafo(doc, WD_ALIGN_PARAGRAPH.JUSTIFY)
    p.paragraph_format.first_line_indent = Cm(1.25)
    adicionar_runs_estilizados(p, texto)
    return p


def paragrafo_verde(doc, texto):
    p = novo_paragrafo(doc)
    p.paragraph_format.first_line_indent = None
    adicionar_runs_estilizados(p, texto)
    return p


def titulo(doc, texto, nivel):
    if nivel == 1:
        p = novo_paragrafo(doc, WD_ALIGN_PARAGRAPH.CENTER)
        adicionar_runs_estilizados(p, texto, tamanho=14, negrito=True)
    elif nivel == 2:
        p = novo_paragrafo(doc, WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.space_before = Pt(12)
        adicionar_runs_estilizados(p, texto, tamanho=12, negrito=True)
    else:
        p = novo_paragrafo(doc, WD_ALIGN_PARAGRAPH.LEFT)
        adicionar_runs_estilizados(p, texto, tamanho=12, negrito=True, italico=True)
    p.paragraph_format.keep_with_next = True
    return p


def quebra_pagina(doc):
    doc.add_page_break()


def sombrear_celula(celula, cor="D9D9D9"):
    tcpr = celula._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), cor)
    tcpr.append(shd)


def limpar_espacos(texto):
    return re.sub(r"\s+", " ", texto.strip())


def tabela_md(doc, linhas):
    dados = []
    for linha in linhas:
        conteudo = linha.strip().strip("|")
        celulas = [limpar_espacos(c) for c in conteudo.split("|")]
        if re.match(r"^[\s:\-]+$", "".join(celulas)):
            continue
        dados.append(celulas)

    if not dados:
        return

    colunas = max(len(r) for r in dados)
    tabela = doc.add_table(rows=0, cols=colunas)
    tabela.style = "Table Grid"
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, linha in enumerate(dados):
        linha_ok = linha + [""] * (colunas - len(linha))
        celulas_docx = tabela.add_row().cells
        for j, texto in enumerate(linha_ok):
            celula = celulas_docx[j]
            p = celula.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(2)
            if i == 0:
                adicionar_runs_estilizados(p, texto, tamanho=12, negrito=True)
                sombrear_celula(celula)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                adicionar_runs_estilizados(p, texto, tamanho=12)

    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def configurar_base(doc):
    estilos = doc.styles
    normal = estilos["Normal"]
    normal.font.name = FONTE
    normal.font.size = Pt(12)
    normal.font.color.rgb = PRETO
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), FONTE)
    rfonts.set(qn("w:hAnsi"), FONTE)
    rfonts.set(qn("w:eastAsia"), FONTE)

    secao = doc.sections[0]
    secao.top_margin = Cm(3)
    secao.bottom_margin = Cm(2)
    secao.left_margin = Cm(3)
    secao.right_margin = Cm(2)


def montar_capa(doc, dados, titulo_projeto):
    def centralizada(texto, tamanho=12, negrito=False, italico=False, antes=0, depois=0):
        p = novo_paragrafo(doc, WD_ALIGN_PARAGRAPH.CENTER)
        p.paragraph_format.space_before = Pt(antes)
        p.paragraph_format.space_after = Pt(depois)
        adicionar_runs_estilizados(p, texto, tamanho=tamanho, negrito=negrito, italico=italico)

    centralizada(dados.get("instituicao", "SERVIÇO NACIONAL DE APRENDIZAGEM INDUSTRIAL - SENAI"),
                 tamanho=13, negrito=True, depois=24)
    centralizada(f"Curso: {dados.get('curso', 'Desenvolvimento de Sistemas')}", tamanho=12, depois=12)
    centralizada("Trabalho de Conclusão de Curso", tamanho=12, negrito=True, depois=60)
    centralizada(titulo_projeto, tamanho=16, negrito=True, depois=60)
    if dados.get("integrantes"):
        centralizada(f"Integrantes: {dados['integrantes']}", depois=6)
    if dados.get("orientador"):
        centralizada(f"Professor Orientador: {dados['orientador']}", depois=6)
    if dados.get("cidade"):
        centralizada(dados["cidade"], depois=6)
    if dados.get("ano"):
        centralizada(dados["ano"], depois=6)
    quebra_pagina(doc)


def montar_sumario(doc):
    p = novo_paragrafo(doc, WD_ALIGN_PARAGRAPH.CENTER)
    adicionar_runs_estilizados(p, "SUMÁRIO", tamanho=14, negrito=True)
    p.paragraph_format.space_after = Pt(18)

    itens = [
        "1. Introdução", "2. Objetivos", "3. Descrição do Problema",
        "4. Tecnologias Utilizadas", "5. Levantamento de Requisitos",
        "6. Modelagem do Sistema", "7. Banco de Dados",
        "8. Desenvolvimento do Back-end", "9. Desenvolvimento do Front-end",
        "10. Funcionalidades do Sistema", "11. Testes Realizados",
        "12. Dificuldades Encontradas", "13. Melhorias Futuras",
        "14. Conclusão", "15. Referências", "Anexos"
    ]
    for item in itens:
        p = novo_paragrafo(doc, WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.first_line_indent = Cm(0)
        adicionar_runs_estilizados(p, item)
    quebra_pagina(doc)


def extrair(capa_linhas):
    dados = {}
    for texto in capa_linhas:
        limpo = texto.strip().strip("*")
        if re.fullmatch(r"\d{4}", limpo):
            dados["ano"] = limpo
            continue
        if ":" in limpo and not limpo.startswith("SERVIÇO") and not limpo.startswith("Serviço"):
            rotulo, valor = limpo.split(":", 1)
            chave = {
                "Curso": "curso",
                "Título": "titulo",
                "Integrantes": "integrantes",
                "Professor Orientador": "orientador",
                "Cidade": "cidade",
            }.get(rotulo.strip(), None)
            if chave:
                dados[chave] = valor.strip()
    return dados


def principal():
    with open(MD_PATH, encoding="utf-8") as f:
        linhas = f.read().splitlines()

    doc = Document()
    configurar_base(doc)

    # Sessões
    indice_capa = next(i for i, l in enumerate(linhas) if l.strip() == "## CAPA")
    indice_sumario = next(i for i, l in enumerate(linhas) if l.strip() == "## SUMÁRIO")
    indice_intro = next(i for i, l in enumerate(linhas) if l.strip() == "## 1. INTRODUÇÃO")

    info_capa = extrair(linhas[indice_capa + 1: indice_sumario])
    titulo_projeto = info_capa.get("titulo",
                                   "LFG - Look For Group: Plataforma Web para Conexão de Jogadores")

    montar_capa(doc, info_capa, titulo_projeto)
    montar_sumario(doc)

    # Corpo do documento
    in_code = False
    linhas_tabela = []
    espera_fim_linha = False

    i = indice_intro
    while i < len(linhas):
        linha_atual = linhas[i]
        bruto = linha_atual.rstrip()

        if in_code:
            if bruto.strip().startswith("```"):
                in_code = False
                i += 1
                continue
            p = novo_paragrafo(doc, WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(bruto)
            estilo_fonte(run, tamanho=10, fonte="Consolas")
            i += 1
            continue

        if not bruto.strip():
            i += 1
            continue

        if bruto.strip().startswith("```"):
            in_code = True
            i += 1
            continue

        if bruto.startswith("|"):
            linhas_tabela.append(bruto)
            i += 1
            continue
        if linhas_tabela:
            tabela_md(doc, linhas_tabela)
            linhas_tabela = []

        limite = bruto.strip()
        if limite.startswith("---"):
            i += 1
            continue

        if limite.startswith("## #"):
            i += 1
            continue

        if limite.startswith("### "):
            titulo(doc, limite[4:], 3)
            i += 1
            continue

        if limite.startswith("## "):
            titulo(doc, limite[3:], 2)
            i += 1
            continue

        if limite.startswith("# "):
            titulo(doc, limite[2:], 1)
            i += 1
            continue

        if limite.startswith("> "):
            texto = limite[2:]
            p = novo_paragrafo(doc, WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.first_line_indent = Cm(0)
            adicionar_runs_estilizados(p, texto, tamanho=12, italico=True)
            i += 1
            continue

        if re.match(r"^- ", limite):
            texto = limite[2:]
            p = novo_paragrafo(doc, WD_ALIGN_PARAGRAPH.JUSTIFY)
            p.paragraph_format.left_indent = Cm(1.0)
            run = p.add_run("\u2022  ")
            estilo_fonte(run)
            adicionar_runs_estilizados(p, texto)
            i += 1
            continue

        if re.match(r"^\d+\. ", limite):
            texto = re.sub(r"^\d+\. ", "", limite)
            p = novo_paragrafo(doc, WD_ALIGN_PARAGRAPH.JUSTIFY)
            p.paragraph_format.left_indent = Cm(1.0)
            adicionar_runs_estilizados(p, texto)
            i += 1
            continue

        if limite.startswith("**Tabela"):
            p = novo_paragrafo(doc, WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            adicionar_runs_estilizados(p, texto=limite, negrito=False)
            for run in p.runs:
                run.italic = True
            i += 1
            continue

        if limite.startswith("**"):
            p = novo_paragrafo(doc, WD_ALIGN_PARAGRAPH.LEFT)
            adicionar_runs_estilizados(p, limite)
            i += 1
            continue

        paragrafo_corpo(doc, limite)
        i += 1

    if linhas_tabela:
        tabela_md(doc, linhas_tabela)

    doc.save(OUT_PATH)
    print(f"Documento gerado: {OUT_PATH}")


if __name__ == "__main__":
    principal()